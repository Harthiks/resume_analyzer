"""
File parsing service — extracts text from PDF and DOCX files.
"""
import io
from typing import Optional

async def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF file bytes using pdfplumber."""
    text = ""
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        # Fallback to PyPDF2
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                text += page.extract_text() or ""
        except Exception as e2:
            print(f"PDF parse error: {e2}")
    return text.strip()

async def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX file bytes using python-docx."""
    text = ""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + " "
                text += "\n"
    except Exception as e:
        print(f"DOCX parse error: {e}")
    return text.strip()

async def extract_text(file_bytes: bytes, filename: str) -> str:
    """Auto-detect file type and extract text."""
    fname_lower = filename.lower()
    if fname_lower.endswith(".pdf"):
        return await extract_text_from_pdf(file_bytes)
    elif fname_lower.endswith(".docx") or fname_lower.endswith(".doc"):
        return await extract_text_from_docx(file_bytes)
    elif fname_lower.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")
    else:
        # Try PDF first, then DOCX
        text = await extract_text_from_pdf(file_bytes)
        if not text:
            text = await extract_text_from_docx(file_bytes)
        return text
